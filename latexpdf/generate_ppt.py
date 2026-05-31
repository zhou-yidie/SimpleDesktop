# -*- coding: utf-8 -*-
import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    # 初始化 16:9 宽屏幻灯片
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 颜色系统定义
    NAVY = RGBColor(16, 44, 87)       # #102C57 深海军蓝 (主色)
    AMBER = RGBColor(228, 130, 87)    # #E48257 琥珀金/暖色 (强调色)
    WHITE = RGBColor(255, 255, 255)   # #FFFFFF 纯白
    LIGHT_BG = RGBColor(243, 244, 246)# #F3F4F6 浅灰色背景
    DARK_TEXT = RGBColor(31, 41, 55)   # #1F2937 深碳灰
    GRAY_TEXT = RGBColor(107, 114, 128)# #6B7280 次要灰

    # 字体定义
    FONT_TITLE = "Microsoft YaHei"
    FONT_BODY = "SimSun"

    # 辅助函数：设置文本框的边距为 0，防止排版偏移
    def zero_margins(tf):
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)
        tf.margin_left = Inches(0)
        tf.margin_right = Inches(0)

    # 辅助函数：在幻灯片上绘制背景颜色
    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # 辅助函数：创建一个通用的有顶部标题栏的内容幻灯片
    def create_content_slide(title_text):
        slide = prs.slides.add_slide(prs.slide_layouts[6]) # 空白版式
        set_slide_background(slide, LIGHT_BG)

        # 绘制顶部深蓝标题装饰条
        header_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1)
        )
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = NAVY
        header_shape.line.color.rgb = NAVY

        # 绘制顶部标题下方的琥珀金装饰细线
        line_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.1), Inches(13.333), Inches(0.06)
        )
        line_shape.fill.solid()
        line_shape.fill.fore_color.rgb = AMBER
        line_shape.line.color.rgb = AMBER

        # 添加标题文字
        tx_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.2), Inches(12.0), Inches(0.7))
        tf = tx_box.text_frame
        zero_margins(tf)
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_TITLE
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.LEFT
        
        return slide

    # 辅助函数：安全地添加图片，如果不存在则画一个带文字的占位符
    def safe_add_picture(slide, img_name, left, top, width=None, height=None):
        img_path = os.path.join("figures", img_name)
        if os.path.exists(img_path):
            try:
                if width and height:
                    slide.shapes.add_picture(img_path, left, top, width=width, height=height)
                elif width:
                    slide.shapes.add_picture(img_path, left, top, width=width)
                elif height:
                    slide.shapes.add_picture(img_path, left, top, height=height)
                else:
                    slide.shapes.add_picture(img_path, left, top)
            except Exception as e:
                # 异常处理：绘制红色框提示
                box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width or Inches(4), height or Inches(3))
                box.fill.solid()
                box.fill.fore_color.rgb = RGBColor(254, 226, 226)
                box.line.color.rgb = AMBER
                tf = box.text_frame
                tf.text = f"图片加载失败:\n{img_name}\n{str(e)}"
                tf.paragraphs[0].font.size = Pt(12)
                tf.paragraphs[0].font.color.rgb = AMBER
        else:
            # 占位符处理
            box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width or Inches(4.5), height or Inches(3.5))
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor(243, 244, 246)
            box.line.color.rgb = GRAY_TEXT
            tf = box.text_frame
            tf.text = f"【暂无配图】\n{img_name}"
            p = tf.paragraphs[0]
            p.font.name = FONT_TITLE
            p.font.size = Pt(14)
            p.font.color.rgb = GRAY_TEXT
            p.alignment = PP_ALIGN.CENTER

    # -------------------------------------------------------------
    # SLIDE 1: 封面 (Sleek 深蓝背景)
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide1, NAVY)

    # 封面左侧琥珀金装饰竖条
    bar_shape = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.3), Inches(7.5)
    )
    bar_shape.fill.solid()
    bar_shape.fill.fore_color.rgb = AMBER
    bar_shape.line.color.rgb = AMBER

    # 主副标题框
    title_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(2.5))
    tf1 = title_box.text_frame
    zero_margins(tf1)
    tf1.word_wrap = True

    p_title = tf1.paragraphs[0]
    p_title.text = "基于微内核与无障碍辅助技术的\n适老化智能桌面系统设计与实现"
    p_title.font.name = FONT_TITLE
    p_title.font.size = Pt(40)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.alignment = PP_ALIGN.LEFT

    p_subtitle = tf1.add_paragraph()
    p_subtitle.text = "—— 毕业设计论文答辩汇报 (SimpleDesktop)"
    p_subtitle.font.name = FONT_TITLE
    p_subtitle.font.size = Pt(22)
    p_subtitle.font.color.rgb = AMBER
    p_subtitle.space_before = Pt(20)
    p_subtitle.alignment = PP_ALIGN.LEFT

    # 作者及导师信息框
    info_box = slide1.shapes.add_textbox(Inches(1.2), Inches(4.8), Inches(10.0), Inches(1.8))
    tf_info = info_box.text_frame
    zero_margins(tf_info)
    tf_info.word_wrap = True

    infos = [
        "汇报人：[您的姓名]     专业：通信工程 / 软件工程",
        "指导老师：[导师姓名]     学院：信息与通信工程学院",
        "答辩时间：2026 年 6 月"
    ]
    for i, info in enumerate(infos):
        p = tf_info.add_paragraph() if i > 0 else tf_info.paragraphs[0]
        p.text = info
        p.font.name = FONT_TITLE
        p.font.size = Pt(16)
        p.font.color.rgb = WHITE
        p.space_before = Pt(8) if i > 0 else Pt(0)
        p.alignment = PP_ALIGN.LEFT


    # -------------------------------------------------------------
    # SLIDE 2: 课题背景与现有痛点 (配人口老龄化趋势图)
    # -------------------------------------------------------------
    slide2 = create_content_slide("01. 课题背景与现有痛点 (Context)")
    
    # 左侧文字卡片 (极简精炼)
    left_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.3))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = WHITE
    left_card.line.color.rgb = RGBColor(229, 231, 235)
    
    tf_lc = left_card.text_frame
    zero_margins(tf_lc)
    tf_lc.word_wrap = True
    
    p = tf_lc.paragraphs[0]
    p.text = "📊 痛点剖析与适老需求"
    p.font.name = FONT_TITLE
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_before = Pt(10)
    
    p2 = tf_lc.add_paragraph()
    p2.text = "• 老龄化形势严峻：\n  2025年全国老年人口达3.21亿(占比22.86%)，老龄化进程提速。\n• 现有极简模式断层：\n  传统厂商极简模式仅停留在桌面层（大字号图标），一旦进入第三方应用立刻失效——“系统层优化充分、应用层穿透不足”。\n• 通信与网络壁垒：\n  微信通话操作繁多（需7步）；断网后缺乏自愈能力，误触即失联。"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(15)
    p2.font.color.rgb = DARK_TEXT
    p2.space_before = Pt(20)
    p2.line_spacing = 1.3

    # 右侧配图：中国老年人口趋势图
    right_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = WHITE
    right_card.line.color.rgb = RGBColor(229, 231, 235)
    tf_rc = right_card.text_frame
    zero_margins(tf_rc)
    tf_rc.word_wrap = True
    p = tf_rc.paragraphs[0]
    p.text = "📈 老年人口及占比变动趋势图"
    p.font.name = FONT_TITLE
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_before = Pt(10)
    p.alignment = PP_ALIGN.CENTER

    # 嵌入趋势图 fig1_1_aging_population_trend.png
    safe_add_picture(slide2, "fig1_1_aging_population_trend.png", Inches(7.1), Inches(2.3), width=Inches(5.3))


    # -------------------------------------------------------------
    # SLIDE 3: 核心重点 - 研究意义与架构对比图
    # -------------------------------------------------------------
    slide3 = create_content_slide("02. 【核心亮点】研究意义与架构对比 (Significance)")
    
    # 左侧文字卡片：三维立体研究价值 (高度提炼)
    left_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.3))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = WHITE
    left_card.line.color.rgb = RGBColor(229, 231, 235)
    tf_lc = left_card.text_frame
    zero_margins(tf_lc)
    tf_lc.word_wrap = True
    
    p = tf_lc.paragraphs[0]
    p.text = "🌟 三维立体研究与工程价值"
    p.font.name = FONT_TITLE
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_before = Pt(10)
    
    p2 = tf_lc.add_paragraph()
    p2.text = "• 社会价值：消弭“数字鸿沟”，将多步繁琐沟通转化为一键直达，保障独居老人亲情维系与急救。\n• 工程价值：设计宿主+总线+插件的三层微内核，实现业务故障隔离与低配终端动态保活。\n• 学术与创新价值：打破传统“止步于Launcher桌面”的限制，首创“穿透式无障碍自动代理”理论与多级鲁棒控制算法。"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(15)
    p2.font.color.rgb = DARK_TEXT
    p2.space_before = Pt(20)
    p2.line_spacing = 1.35

    # 右侧卡片：架构对比图 (展示 Launcher 局限与本系统的穿透式设计)
    right_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = WHITE
    right_card.line.color.rgb = RGBColor(229, 231, 235)
    tf_rc = right_card.text_frame
    zero_margins(tf_rc)
    tf_rc.word_wrap = True
    p = tf_rc.paragraphs[0]
    p.text = "🔄 传统改造 vs 穿透式无障碍架构对比"
    p.font.name = FONT_TITLE
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_before = Pt(10)
    p.alignment = PP_ALIGN.CENTER

    # 嵌入架构对比图 fig1_3_architecture_compare.png
    safe_add_picture(slide3, "fig1_3_architecture_compare.png", Inches(7.0), Inches(2.3), width=Inches(5.5))


    # -------------------------------------------------------------
    # SLIDE 4: 系统总体架构设计 (配三层微内核架构模块图)
    # -------------------------------------------------------------
    slide4 = create_content_slide("03. 系统总体架构与技术解耦 (Architecture)")
    
    # 左侧文字卡片
    left_c = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.3))
    left_c.fill.solid()
    left_c.fill.fore_color.rgb = WHITE
    left_c.line.color.rgb = RGBColor(229, 231, 235)
    tf_lc = left_c.text_frame
    zero_margins(tf_lc)
    tf_lc.word_wrap = True
    
    p = tf_lc.paragraphs[0]
    p.text = "🏗️ 三层微内核架构拓扑"
    p.font.name = FONT_TITLE
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_before = Pt(10)
    
    p2 = tf_lc.add_paragraph()
    p2.text = "• 宿主层 (Host)：系统稳定核心。提供统一生命周期、Dagger-Hilt依赖注入、DataStore本地加密存储。\n• 总线层 (Bus)：插件管理枢纽。基于消息总线发布/订阅模型，隔离各插件，杜绝级联崩溃。\n• 插件层 (Plugins)：高频业务高度插件化。主屏、联系人、自愈网络及语音助手热插拔，可按需精简。"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(15)
    p2.font.color.rgb = DARK_TEXT
    p2.space_before = Pt(20)
    p2.line_spacing = 1.35

    # 右侧配图：系统核心模块架构图
    right_c = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3))
    right_c.fill.solid()
    right_c.fill.fore_color.rgb = WHITE
    right_c.line.color.rgb = RGBColor(229, 231, 235)
    tf_rc = right_c.text_frame
    zero_margins(tf_rc)
    tf_rc.word_wrap = True
    p = tf_rc.paragraphs[0]
    p.text = "🧩 SimpleDesktop 系统核心功能模块结构"
    p.font.name = FONT_TITLE
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_before = Pt(10)
    p.alignment = PP_ALIGN.CENTER

    # 嵌入模块架构图 fig3_1_system_modules.png
    safe_add_picture(slide4, "fig3_1_system_modules.png", Inches(7.0), Inches(2.2), width=Inches(5.5))


    # -------------------------------------------------------------
    # SLIDE 5: 核心重点 - 微信自动通话（七阶段原子状态机与流转图）
    # -------------------------------------------------------------
    slide5 = create_content_slide("04. 【核心亮点】微信自动通话（七阶段原子状态机）")
    
    # 左侧文字卡片：状态机详解
    left_c = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.3))
    left_c.fill.solid()
    left_c.fill.fore_color.rgb = WHITE
    left_c.line.color.rgb = RGBColor(229, 231, 235)
    tf_lc = left_c.text_frame
    zero_margins(tf_lc)
    tf_lc.word_wrap = True
    
    p = tf_lc.paragraphs[0]
    p.text = "🔄 离散原子流转与多级鲁棒控制"
    p.font.name = FONT_TITLE
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_before = Pt(10)
    
    p2 = tf_lc.add_paragraph()
    p2.text = "• 七阶段流程离散化：\n  SEARCH $\\rightarrow$ USER $\\rightarrow$ CHAT $\\rightarrow$ PANEL $\\rightarrow$ VIDEO $\\rightarrow$ CONFIRM $\\rightarrow$ CALLING，实现拨视频一点击直达。\n• 强容错鲁棒技术：\n  - 自回归机制：阶段流转前/后校验界面身份，如被卡顿或误触中断，模拟返回键重置触发自回归，成功率达 99.1%。\n  - 双通道定位：Resource-ID定位+层级拓扑模糊检索，完美抵御微信版本更新混淆。\n  - 焦点退避：自动识别来电或广告干扰并退避，5秒心跳锁死防发热。"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(14)
    p2.font.color.rgb = DARK_TEXT
    p2.space_before = Pt(15)
    p2.line_spacing = 1.3

    # 右侧：七阶段状态机流转图
    right_c = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3))
    right_c.fill.solid()
    right_c.fill.fore_color.rgb = WHITE
    right_c.line.color.rgb = RGBColor(229, 231, 235)
    tf_rc = right_c.text_frame
    zero_margins(tf_rc)
    tf_rc.word_wrap = True
    p = tf_rc.paragraphs[0]
    p.text = "🎬 微信一键直拨七阶段状态机流程图"
    p.font.name = FONT_TITLE
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_before = Pt(10)
    p.alignment = PP_ALIGN.CENTER

    # 嵌入状态机流转图 wechat_state.png
    safe_add_picture(slide5, "wechat_state.png", Inches(7.4), Inches(2.2), width=Inches(4.7))


    # -------------------------------------------------------------
    # SLIDE 6: 网络主动守护与自愈机制 (配守护与自愈闭环流程图)
    # -------------------------------------------------------------
    slide6 = create_content_slide("05. 网络主动守护与自愈机制 (WiFi Netguard)")
    
    # 左侧文字卡片
    left_c = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.3))
    left_c.fill.solid()
    left_c.fill.fore_color.rgb = WHITE
    left_c.line.color.rgb = RGBColor(229, 231, 235)
    tf_lc = left_c.text_frame
    zero_margins(tf_lc)
    tf_lc.word_wrap = True
    
    p = tf_lc.paragraphs[0]
    p.text = "🔋 监测保活与无障碍自愈闭环"
    p.font.name = FONT_TITLE
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_before = Pt(10)
    
    p2 = tf_lc.add_paragraph()
    p2.text = "• 常驻前台保活：挂载低优先级静默通知防止守护服务被后台杀死；通过底层广播捕获断网瞬间。\n• Android 11-16 无障碍自愈：克服Android 10以上禁封直接调用开启 API 限制。\n  断网瞬间唤起无障碍辅助 $\\rightarrow$ 跳转系统原生WLAN设置页 $\\rightarrow$ 自动遍历定位并模拟点击WiFi开关 $\\rightarrow$ 成功后模拟Home键返回。全程仅需 4~5 秒且老人无感。\n• 多级防御宽限期：设置 10 秒宽限期过滤短暂抖动；若物理故障失败则触发 0.9 倍慢速 TTS 语音播报及引导卡片。"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(14)
    p2.font.color.rgb = DARK_TEXT
    p2.space_before = Pt(15)
    p2.line_spacing = 1.3

    # 右侧配图：网络守护自愈机制流程图
    right_c = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3))
    right_c.fill.solid()
    right_c.fill.fore_color.rgb = WHITE
    right_c.line.color.rgb = RGBColor(229, 231, 235)
    tf_rc = right_c.text_frame
    zero_margins(tf_rc)
    tf_rc.word_wrap = True
    p = tf_rc.paragraphs[0]
    p.text = "🛡️ WiFi 网络守护与自愈机制流程"
    p.font.name = FONT_TITLE
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_before = Pt(10)
    p.alignment = PP_ALIGN.CENTER

    # 嵌入流程图 wifi_daemon.png
    safe_add_picture(slide6, "wifi_daemon.png", Inches(7.0), Inches(2.2), width=Inches(5.5))


    # -------------------------------------------------------------
    # SLIDE 7: 自然语言语音交互与容错 (配语音交互流程图)
    # -------------------------------------------------------------
    slide7 = create_content_slide("06. 多模态辅助：语音助手与语义容错 (Voice)")
    
    # 左侧文字卡片
    left_c = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.3))
    left_c.fill.solid()
    left_c.fill.fore_color.rgb = WHITE
    left_c.line.color.rgb = RGBColor(229, 231, 235)
    tf_lc = left_c.text_frame
    zero_margins(tf_lc)
    tf_lc.word_wrap = True
    
    p = tf_lc.paragraphs[0]
    p.text = "🗣️ 本地极轻量双层解析与容错"
    p.font.name = FONT_TITLE
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_before = Pt(10)
    
    p2 = tf_lc.add_paragraph()
    p2.text = "• 双层轻量语义解析：\n  - 第一层：关键词槽位提取，专为中低配机型设计，不依赖云 API。\n  - 第二层：编辑距离 (Levenshtein) 容错。允许发音模糊或 1 字误差（如“打电活给张三”）自动命中意图。\n• 四层联系人阶梯检索：\n  完全匹配 $\\rightarrow$ 包含匹配 $\\rightarrow$ 拼音首字母匹配 $\\rightarrow$ 编辑距离匹配。置信度综合计分检索，保证呼叫无偏差。\n• 听觉自适应 TTS：反向播报语速默认下调至正常的 0.9 倍，语调舒缓贴心。"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(14.5)
    p2.font.color.rgb = DARK_TEXT
    p2.space_before = Pt(15)
    p2.line_spacing = 1.3

    # 右侧配图：语音交互流程图
    right_c = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3))
    right_c.fill.solid()
    right_c.fill.fore_color.rgb = WHITE
    right_c.line.color.rgb = RGBColor(229, 231, 235)
    tf_rc = right_c.text_frame
    zero_margins(tf_rc)
    tf_rc.word_wrap = True
    p = tf_rc.paragraphs[0]
    p.text = "🎙️ 语音辅助模块工作流程"
    p.font.name = FONT_TITLE
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_before = Pt(10)
    p.alignment = PP_ALIGN.CENTER

    # 嵌入语音流程图 voice_flow.png
    safe_add_picture(slide7, "voice_flow.png", Inches(7.0), Inches(2.2), width=Inches(5.5))


    # -------------------------------------------------------------
    # SLIDE 8: 界面呈现与核心功能展示 (配实际系统运行双截图)
    # -------------------------------------------------------------
    slide8 = create_content_slide("07. 系统界面呈现与视觉设计 (Interface)")
    
    # 左侧文字卡片
    left_c = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.3))
    left_c.fill.solid()
    left_c.fill.fore_color.rgb = WHITE
    left_c.line.color.rgb = RGBColor(229, 231, 235)
    tf_lc = left_c.text_frame
    zero_margins(tf_lc)
    tf_lc.word_wrap = True
    
    p = tf_lc.paragraphs[0]
    p.text = "🎨 Jetpack Compose 适老交互"
    p.font.name = FONT_TITLE
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_before = Pt(10)
    
    p2 = tf_lc.add_paragraph()
    p2.text = "• Compose 声明式自适应：全系统基于严格 UI Tokens 色板和触摸热区边距开发，多分辨率无变形错位。\n• 响应式字号无感变大：\n  中/大/超大 3 档字号在 80 毫秒内重组渲染，决无控件重叠或文本溢出。\n• WCAG AAA 高对比度主题：一键切换纯黑底、白字、鲜橙亮边的高反差模式，显著改善老年弱视认读力。\n• 点击防连击过滤：\n  多重事件状态锁定，秒级内重复连击自动过滤合并，防范状态机崩溃。"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(14)
    p2.font.color.rgb = DARK_TEXT
    p2.space_before = Pt(15)
    p2.line_spacing = 1.3

    # 右侧：展示实机界面大图 (双图并列)
    right_c = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3))
    right_c.fill.solid()
    right_c.fill.fore_color.rgb = WHITE
    right_c.line.color.rgb = RGBColor(229, 231, 235)
    tf_rc = right_c.text_frame
    zero_margins(tf_rc)
    tf_rc.word_wrap = True
    p = tf_rc.paragraphs[0]
    p.text = "📱 桌面主界面及高反差界面实拍"
    p.font.name = FONT_TITLE
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_before = Pt(10)
    p.alignment = PP_ALIGN.CENTER

    # 嵌入主界面演示图 demo1.png 和 demo2.png
    safe_add_picture(slide8, "demo1.png", Inches(7.4), Inches(2.2), width=Inches(2.2))
    safe_add_picture(slide8, "demo2.png", Inches(9.9), Inches(2.2), width=Inches(2.2))


    # -------------------------------------------------------------
    # SLIDE 9: 兼容性测试与真机环境验证 (保留精细表格)
    # -------------------------------------------------------------
    slide9 = create_content_slide("08. 兼容性测试与真机环境验证 (Testing)")
    
    # 顶部加粗测试结论
    intro_box = slide9.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12.0), Inches(0.8))
    tf_i = intro_box.text_frame
    zero_margins(tf_i)
    tf_i.word_wrap = True
    p_i = tf_i.paragraphs[0]
    p_i.text = "📊 本系统在跨芯片平台、屏幕分辨率及不同系统版本（Android 11至16，以及 HarmonyOS NEXT）上展开了严苛验证。"
    p_i.font.name = FONT_TITLE
    p_i.font.size = Pt(16)
    p_i.font.bold = True
    p_i.font.color.rgb = NAVY

    # 创建兼容性设备表格 (6款代表性真机)
    table_shape = slide9.shapes.add_table(7, 3, Inches(0.6), Inches(2.2), Inches(12.133), Inches(4.5))
    table = table_shape.table
    
    # 设置列宽
    table.columns[0].width = Inches(1.0)
    table.columns[1].width = Inches(2.5)
    table.columns[2].width = Inches(8.633)
    
    # 表头数据
    headers = ["序号", "测试设备", "关键系统参数与兼容表现"]
    for col_idx, header_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = header_text
        p.font.name = FONT_TITLE
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        
    # 表格行数据 (6款设备)
    devices_data = [
        ("1", "一加 13", "骁龙8至尊版 / 12GB RAM / 1440×3168 QHD+ / Android 16 (ColorOS 16)\n表现：超高分辨率矢量图清晰自适应；Android 16 增强无障碍安全策略完美通过"),
        ("2", "荣耀 WinRT", "骁龙8至尊版 / 16GB RAM / 1272×2800 1.5K / Android 15 (MagicOS 10.0)\n表现：MagicOS 后台自启动机制稳定保活，通知授权及前台常驻功能稳定"),
        ("3", "一加 Ace", "天玑8100-MAX / 12GB RAM / 1080×2412 FHD+ / Android 12 (ColorOS 12.1)\n表现：主流千元机主流配置分辨率显示端正，微信视频自愈快速成功"),
        ("4", "iQOO 8", "高通骁龙888 / 12GB RAM / 1080×2376 FHD+ / Android 11 (OriginOS 1.0)\n表现：老旧基准环境无障碍自愈耗时4.8秒，各项前后台运行指标完美适配"),
        ("5", "荣耀 70", "高通骁龙778G Plus / 8GB RAM / 1080×2400 FHD+ / Android 12 (MagicUI 6.1)\n表现：中低配机型后台保活白名单一键跳转流畅，前台无障碍执行稳定可靠"),
        ("6", "华为 Pura 80 Pro", "麒麟9020 / 12GB RAM / 1260×2844 1.5K / HarmonyOS 5.1 (HarmonyOS NEXT)\n表现：支持鸿蒙无障碍全新动态权限机制，通过兼容层在纯血鸿蒙上完美运转")
    ]
    
    for row_idx, data in enumerate(devices_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.name = FONT_TITLE if col_idx < 2 else FONT_BODY
            p.font.size = Pt(12)
            p.font.color.rgb = DARK_TEXT
            p.alignment = PP_ALIGN.CENTER if col_idx < 2 else PP_ALIGN.LEFT


    # -------------------------------------------------------------
    # SLIDE 10: 性能实测与结论展望 (配冷热启动与切换加载耗时图)
    # -------------------------------------------------------------
    slide10 = create_content_slide("09. 性能实测、资源开销与结论展望 (Result)")
    
    # 左侧文字卡片 (高度提炼)
    left_c = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.3))
    left_c.fill.solid()
    left_c.fill.fore_color.rgb = WHITE
    left_c.line.color.rgb = RGBColor(229, 231, 235)
    tf_lc = left_c.text_frame
    zero_margins(tf_lc)
    tf_lc.word_wrap = True
    
    p = tf_lc.paragraphs[0]
    p.text = "📈 8小时压力实测与总结展望"
    p.font.name = FONT_TITLE
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_before = Pt(10)
    
    p2 = tf_lc.add_paragraph()
    p2.text = "• 强悍实测性能：\n  - 满帧渲染：列表滑动稳定在 60.1 fps 满帧流畅。\n  - 极低消耗：8小时常驻内存仅 95~115 MB；空闲 CPU 占用 < 1%，密集操作短时达 8%~12% 快速回落；连续 8 小时整体耗电仅 3%~5%。\n  - 闪电响应：字号/主题切换在 80~90 毫秒内，冷启动 1.15 秒。\n• 毕业成果总结：宿主+总线+插件三层微内核架构；七阶段FSM微信一键直拨；WiFi断网自愈与0.9倍TTS防御体系。\n• 未来展望：集成端侧小型大模型，实现自然多轮情感陪伴；接入 Matter 智能家居，打通家用健康血压计插件扩展。"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.color.rgb = DARK_TEXT
    p2.space_before = Pt(15)
    p2.line_spacing = 1.25

    # 右侧配图：页面加载速度耗时图
    right_c = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3))
    right_c.fill.solid()
    right_c.fill.fore_color.rgb = WHITE
    right_c.line.color.rgb = RGBColor(229, 231, 235)
    tf_rc = right_c.text_frame
    zero_margins(tf_rc)
    tf_rc.word_wrap = True
    p = tf_rc.paragraphs[0]
    p.text = "⏱️ 关键场景页面加载与切换耗时统计"
    p.font.name = FONT_TITLE
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_before = Pt(10)
    p.alignment = PP_ALIGN.CENTER

    # 嵌入耗时图 fig4_load_speed.png
    safe_add_picture(slide10, "fig4_load_speed.png", Inches(7.0), Inches(2.2), width=Inches(5.5))


    # -------------------------------------------------------------
    # SLIDE 11: 致谢 (深海军蓝背景)
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide11, NAVY)

    # 致谢右侧琥珀金装饰竖条
    bar_shape = slide11.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(13.033), Inches(0), Inches(0.3), Inches(7.5)
    )
    bar_shape.fill.solid()
    bar_shape.fill.fore_color.rgb = AMBER
    bar_shape.line.color.rgb = AMBER

    # 致谢大标题框
    thank_box = slide11.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.0), Inches(3.0))
    tf_t = thank_box.text_frame
    zero_margins(tf_t)
    tf_t.word_wrap = True

    p = tf_t.paragraphs[0]
    p.text = "谢谢各位评委老师！"
    p.font.name = FONT_TITLE
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    p2 = tf_t.add_paragraph()
    p2.text = "敬请各位专家教授批评指正！"
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(28)
    p2.font.color.rgb = AMBER
    p2.space_before = Pt(25)
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf_t.add_paragraph()
    p3.text = "答辩人：[您的姓名]"
    p3.font.name = FONT_TITLE
    p3.font.size = Pt(18)
    p3.font.color.rgb = WHITE
    p3.space_before = Pt(40)
    p3.alignment = PP_ALIGN.CENTER

    # 保存演示文稿
    output_filename = "thesis_presentation_v2.pptx"
    prs.save(output_filename)
    print(f"Presentation successfully created and saved as '{output_filename}'!")

if __name__ == "__main__":
    create_presentation()
